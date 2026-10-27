"""MIS report: Excel workbook whose Summary sheet is driven by live formulas."""
from datetime import datetime

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config import GROUPS, REPORT_DIR  # noqa: F401
from ticketing import FMT

FONT = "Arial"
HEAD_FILL = PatternFill("solid", start_color="1F3864")
HEADERS = ["Ticket ID", "Created", "Category", "Priority", "Summary", "Status", "Assigned Group",
           "SLA Due", "Resolved At", "Resolution (min)", "SLA Met", "Auto-Resolved", "Source"]


def _dt(s):
    return datetime.strptime(s, FMT) if s else None


def build(conn, source="lab", path=None):
    q = "SELECT * FROM tickets" + ("" if source == "all" else " WHERE source=?")
    rows = conn.execute(q, () if source == "all" else (source,)).fetchall()
    path = path or REPORT_DIR / f"MIS_report_{datetime.now():%Y%m%d_%H%M}.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws, summ = wb.active, wb.create_sheet("Summary", 0)
    ws.title = "Tickets"

    ws.append(HEADERS)
    for c in ws[1]:
        c.font, c.fill = Font(name=FONT, bold=True, color="FFFFFF"), HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, r in enumerate(rows, start=2):
        ws.append([r["ticket_id"], _dt(r["created_at"]), r["category"], r["priority"], r["summary"],
                   r["status"], r["assigned_group"], _dt(r["sla_due"]), _dt(r["resolved_at"]),
                   f'=IF(I{i}="","",ROUND((I{i}-B{i})*1440,1))',
                   f'=IF(I{i}="","",IF(I{i}<=H{i},"Yes","No"))',
                   "Yes" if r["auto_resolved"] else "No", r["source"]])
    last = max(2, len(rows) + 1)
    for row in ws.iter_rows(min_row=2, max_row=last):
        for c in row:
            c.font = Font(name=FONT, size=10)
        for idx in (1, 7, 8):
            row[idx].number_format = "yyyy-mm-dd hh:mm" if idx != 6 else "General"
        row[8].number_format = "yyyy-mm-dd hh:mm"
    for col, w in zip("ABCDEFGHIJKLM", [18, 17, 18, 9, 60, 11, 20, 17, 17, 14, 9, 13, 8]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:M{last}"

    # ---- Summary sheet (formulas only; nothing hardcoded) ----
    R = lambda col: f"Tickets!{col}2:{col}{last}"  # noqa: E731
    summ["A1"] = "AutoFix Desk - Service Desk MIS Report"
    summ["A1"].font = Font(name=FONT, size=16, bold=True, color="1F3864")
    summ["A2"] = f"Generated {datetime.now():%Y-%m-%d %H:%M} | Data source filter: {source}"
    summ["A2"].font = Font(name=FONT, size=9, italic=True)
    if source == "demo":
        summ["A3"] = "DEMO DATA - synthetic tickets for testing the report. Do not quote these numbers."
        summ["A3"].font = Font(name=FONT, size=10, bold=True, color="C00000")

    kpis = [
        ("Total tickets", f"=COUNTA({R('A')})", "0"),
        ("Resolved", f'=COUNTIF({R("F")},"Resolved")', "0"),
        ("Escalated", f'=COUNTIF({R("F")},"Escalated")', "0"),
        ("Open", f'=COUNTIF({R("F")},"Open")', "0"),
        ("Auto-resolution rate", '=IFERROR(COUNTIF({0},"Yes")/B5,0)'.format(R("L")), "0.0%"),
        ("SLA met % (resolved tickets)", '=IFERROR(COUNTIF({0},"Yes")/B6,0)'.format(R("K")), "0.0%"),
        ("Escalation rate", "=IFERROR(B7/B5,0)", "0.0%"),
        ("Avg resolution time (min)", f"=IFERROR(AVERAGE({R('J')}),0)", "0.0"),
    ]
    summ["A4"], summ["B4"] = "KPI", "Value"
    for c in (summ["A4"], summ["B4"]):
        c.font, c.fill = Font(name=FONT, bold=True, color="FFFFFF"), HEAD_FILL
    for i, (label, formula, fmt) in enumerate(kpis, start=5):
        summ[f"A{i}"], summ[f"B{i}"] = label, formula
        summ[f"A{i}"].font = summ[f"B{i}"].font = Font(name=FONT, size=10)
        summ[f"B{i}"].number_format = fmt

    cats = sorted({r["category"] for r in rows}) or ["(none)"]
    top = 5 + len(kpis) + 1
    summ[f"A{top}"], summ[f"B{top}"] = "Category", "Tickets"
    for c in (summ[f"A{top}"], summ[f"B{top}"]):
        c.font, c.fill = Font(name=FONT, bold=True, color="FFFFFF"), HEAD_FILL
    for j, cat in enumerate(cats, start=top + 1):
        summ[f"A{j}"] = cat
        summ[f"B{j}"] = f"=COUNTIF({R('C')},A{j})"
        summ[f"A{j}"].font = summ[f"B{j}"].font = Font(name=FONT, size=10)

    chart = BarChart()
    chart.title, chart.y_axis.title, chart.legend = "Tickets by category", "Tickets", None
    chart.add_data(Reference(summ, min_col=2, min_row=top, max_row=top + len(cats)), titles_from_data=True)
    chart.set_categories(Reference(summ, min_col=1, min_row=top + 1, max_row=top + len(cats)))
    chart.height, chart.width = 7.5, 14
    summ.add_chart(chart, f"D{4}")

    note = top + len(cats) + 2
    summ[f"A{note}"] = ("Assumptions: SLA targets come from config.py (P1 60m, P2 240m, P3 480m, P4 1440m) "
                        "- typical L1 values, not a client's real policy.")
    summ[f"A{note}"].font = Font(name=FONT, size=9, italic=True)
    summ.column_dimensions["A"].width = 34
    summ.column_dimensions["B"].width = 12
    wb.save(path)
    return path, len(rows)
